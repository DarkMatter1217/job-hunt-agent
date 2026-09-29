"""Write output/Job_Tracker.xlsx from data/jobs.db (stage 9). No LLM.

Sheets
  Jobs      every open job that passed the filters (+ any job you gave a Status, even after it closes)
  Contacts  referral contacts (filled by the referral stage, Phase 5b)
  Runs      recent runs and any failing sources

Your edits survive: before rewriting, the Status / Notes columns (Jobs) and Outreach status / Last contacted
(Contacts) are read back from the existing file and matched by Job ID / LinkedIn URL.
If the file is open in Excel/Numbers, close it first (or the export writes Job_Tracker_new.xlsx).

    python scripts/export_xlsx.py
"""
import json
import sqlite3
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).parent))
import filters  # noqa: E402
import render  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
DB = ROOT / "data" / "jobs.db"
OUT = ROOT / "output" / "Job_Tracker.xlsx"

JOB_COLS = ["Job ID", "Status", "Notes", "Score", "Company", "Tier", "Title", "Location", "Market", "Term", "Posted",
            "First seen", "Open?", "Why it fits", "Gaps", "Flags", "Apply link", "Jobright link", "Resume", "Cover letter",
            "Application", "Contacts"]
USER_JOB_COLS = ["Status", "Notes"]
STATUSES = ["New", "Applied", "Interview", "Offer", "Rejected", "Skip"]
CONTACT_COLS = ["Job ID", "Company", "Name", "Title", "Connection", "Why this person", "LinkedIn", "Email", "Email source",
                "Person brief", "Hook sources", "Connection note", "Referral message", "Outreach status", "Last contacted"]
USER_CONTACT_COLS = ["Outreach status", "Last contacted"]
OUTREACH = ["Not contacted", "Requested", "Connected", "Replied", "Referred", "No reply"]
WIDTH = {"Job ID": 11, "Status": 11, "Notes": 24, "Score": 7, "Company": 20, "Tier": 5, "Title": 42, "Location": 24,
         "Market": 10, "Term": 11, "Posted": 11, "First seen": 11, "Open?": 8, "Why it fits": 60, "Gaps": 34, "Flags": 28,
         "Apply link": 14, "Jobright link": 14, "Resume": 12, "Cover letter": 12, "Application": 13, "Contacts": 9}
HEADER = PatternFill("solid", fgColor="1F4E79")
NEW_ROW = PatternFill("solid", fgColor="FFF2CC")


def read_user_edits(path, sheet, id_col, cols):
    if not path.exists():
        return {}
    try:
        ws = load_workbook(path, read_only=True)[sheet]
    except Exception:
        return {}
    rows = ws.iter_rows(values_only=True)
    header = list(next(rows, []))
    if id_col not in header:
        return {}
    idx = {c: header.index(c) for c in cols if c in header}
    return {r[header.index(id_col)]: {c: r[i] for c, i in idx.items()} for r in rows if r and r[header.index(id_col)]}


def link(cell, url, text):
    if url:
        cell.value, cell.hyperlink, cell.font = text, url, Font(color="1F4E9A", underline="single")


def term_of(j):
    t = filters.terms_in(j["title"] or "") or filters.terms_in((j["description"] or "")[:2000])
    if not t:
        ht = json.loads(j["extra"] or "{}").get("hire_time")
        return ht or ""
    y, m = sorted(set(t))[0]
    return {1: "Winter", 5: "Summer", 9: "Fall"}.get(m, f"{m:02d}/") + f" {y}"


def main():
    con = sqlite3.connect(DB); con.row_factory = sqlite3.Row
    edits = read_user_edits(OUT, "Jobs", "Job ID", USER_JOB_COLS)
    cedits = read_user_edits(OUT, "Contacts", "LinkedIn", USER_CONTACT_COLS)
    last_run = con.execute("SELECT MAX(run_id) FROM runs").fetchone()[0] or 0
    keep = tuple(k for k, v in edits.items() if v.get("Status") and v["Status"] != "New") or ("",)
    rows = con.execute(f"""SELECT * FROM jobs WHERE (status='open' AND filter_result='pass')
                           OR key IN ({','.join('?' * len(keep))})""", keep).fetchall()
    jobright = {r["key"]: r["url"] for r in con.execute("SELECT key, url FROM sightings WHERE source LIKE 'intern-list:%'")}
    has_contacts = "contacts" in {r[0] for r in con.execute("SELECT name FROM sqlite_master WHERE type='table'")}
    ncontacts = dict(con.execute("SELECT job_key, COUNT(*) FROM contacts GROUP BY job_key").fetchall()) if has_contacts else {}

    def market(j):
        f = j["flags"] or ""
        return next((m for m in ("home", "canada", "remote_unspecified", "india", "us") if f'"market:{m}"' in f), "")

    def sort_key(j):
        return (j["score"] is None, market(j) == "us", -(j["score"] or 0), j["tier"] or 9, j["company"])
    rows = sorted(rows, key=sort_key)

    wb = Workbook()
    ws = wb.active; ws.title = "Jobs"
    ws.append(JOB_COLS)
    for j in rows:
        reason = json.loads(j["score_reason"] or "{}")
        flags = [f["flag"] for f in json.loads(j["flags"] or "[]") if not f["flag"].startswith(("family:", "market:"))]
        out = Path(j["output_dir"]) if j["output_dir"] else None
        user = edits.get(j["key"], {})
        ws.append([j["key"], user.get("Status") or "New", user.get("Notes") or "",
                   j["score"] if j["score"] is not None else ("US: not scored" if market(j) == "us" else "pending"),
                   j["company"], j["tier"], j["title"], (j["location"] or "")[:120], market(j), term_of(j),
                   (filters.parse_posted(j["posted"]) or "") and filters.parse_posted(j["posted"]).isoformat(),
                   (j["first_seen_at"] or "")[:10], "Yes" if j["status"] == "open" else "Closed",
                   reason.get("why", ""), ", ".join(json.loads(j["gaps"] or "[]")), ", ".join(flags),
                   None, None, None, None, j["tailor_status"] or "", ncontacts.get(j["key"], 0)])
        r = ws.max_row
        direct = j["url"] if "jobright.ai" not in (j["url"] or "") else None
        link(ws.cell(r, JOB_COLS.index("Apply link") + 1), direct or jobright.get(j["key"]), "Apply")
        link(ws.cell(r, JOB_COLS.index("Jobright link") + 1), jobright.get(j["key"]), "Jobright")
        if out and out.exists():
            res = next(out.glob(f"{render.doc_prefix()}_Resume_*.pdf"), None)
            cov = next(out.glob(f"{render.doc_prefix()}_Cover_Letter_*.pdf"), None)
            link(ws.cell(r, JOB_COLS.index("Resume") + 1), res and res.as_uri(), "Resume PDF")
            link(ws.cell(r, JOB_COLS.index("Cover letter") + 1), cov and cov.as_uri(), "Letter PDF")
        if j["first_seen_run"] == last_run and last_run > 1:
            for c in ws[r][:3]:
                c.fill = NEW_ROW
    for i, c in enumerate(JOB_COLS, 1):
        cell = ws.cell(1, i); cell.font = Font(bold=True, color="FFFFFF"); cell.fill = HEADER
        ws.column_dimensions[cell.column_letter].width = WIDTH.get(c, 12)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=JOB_COLS[c.column - 1] in ("Why it fits", "Gaps", "Notes"))
    ws.freeze_panes = "E2"; ws.auto_filter.ref = ws.dimensions
    dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=True)
    ws.add_data_validation(dv); dv.add(f"B2:B{max(ws.max_row, 2)}")
    score_col = ws.cell(1, JOB_COLS.index("Score") + 1).column_letter
    ws.conditional_formatting.add(f"{score_col}2:{score_col}{max(ws.max_row, 2)}",
                                  CellIsRule(operator="greaterThanOrEqual", formula=["70"], fill=PatternFill("solid", fgColor="C6EFCE")))

    cs = wb.create_sheet("Contacts"); cs.append(CONTACT_COLS)
    if has_contacts:
        for c in con.execute("SELECT * FROM contacts ORDER BY job_key"):
            u = cedits.get(c["linkedin"], {})
            cs.append([c["job_key"], c["company"], c["name"], c["title"], c["connection"], c["why"], c["linkedin"],
                       c["email"], c["email_source"], c["brief"], c["hook_sources"], c["note_draft"], c["message_draft"],
                       u.get("Outreach status") or "Not contacted", u.get("Last contacted") or ""])
    for i in range(1, len(CONTACT_COLS) + 1):
        cell = cs.cell(1, i); cell.font = Font(bold=True, color="FFFFFF"); cell.fill = HEADER
        cs.column_dimensions[cell.column_letter].width = 22
    dv2 = DataValidation(type="list", formula1='"' + ",".join(OUTREACH) + '"', allow_blank=True)
    cs.add_data_validation(dv2); dv2.add(f"N2:N{max(cs.max_row, 2)}"); cs.freeze_panes = "D2"

    rs = wb.create_sheet("Runs")
    rs.append(["Run", "Started (UTC)", "Fetched", "New", "Open", "Closed", "Sources failed"])
    for r in con.execute("SELECT * FROM runs ORDER BY run_id DESC LIMIT 15"):
        rs.append([r["run_id"], r["started_at"], r["n_fetched"], r["n_new"], r["n_open"], r["n_closed"], r["n_sources_failed"]])
    rs.append([]); rs.append(["Failing sources in the last run"])
    for r in con.execute("SELECT source, company, error FROM board_runs WHERE run_id=? AND ok=0", (last_run,)):
        rs.append([r["source"], r["company"], r["error"]])
    for i in range(1, 8):
        rs.cell(1, i).font = Font(bold=True)
        rs.column_dimensions[rs.cell(1, i).column_letter].width = 18

    OUT.parent.mkdir(parents=True, exist_ok=True)
    target = OUT
    try:
        wb.save(OUT)
    except PermissionError:
        target = OUT.with_name("Job_Tracker_new.xlsx"); wb.save(target)
    scored = sum(1 for j in rows if j["score"] is not None)
    print(json.dumps({"file": str(target), "jobs": len(rows), "scored": scored,
                      "ready_applications": sum(1 for j in rows if j["tailor_status"] == "ready"),
                      "kept_user_edits": len(edits)}))


if __name__ == "__main__":
    main()
